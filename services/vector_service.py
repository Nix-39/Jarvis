"""
VectorService - long-term semantic memory for Jarvis, backed by ChromaDB.

Scope (v1):
    - Semantic search over ALL conversation history (all agents), synced from
      MemoryService. SQLite remains the single source of truth; the vector
      index is derived data that can always be rebuilt.
    - Semantic search over personal documents in data/documents/
      (.txt, .md, .pdf, .docx), organized in category subfolders of any depth.

Design principles:
    - Self-healing sync: anything missed (e.g. Ollama temporarily down) is
      picked up on the next sync. New, changed and deleted documents are
      detected automatically.
    - Changing the embedding model triggers an automatic full re-index,
      because vectors from different models are not comparable.
    - Fail-soft: a failure in long-term memory never breaks an agent reply;
      the agent simply answers without background context.

Explicitly out of scope (v1):
    - OCR for scanned PDFs
    - Reminders / background scheduling (future scheduler can call sync_all())
    - LLM tool-calling for search

Architecture:
    Agent -> VectorService -> (MemoryService, OllamaService.embed, ChromaDB)

Agents only call build_context(). Sync runs:
    - fully at Orchestrator startup,
    - with a small time budget before every search,
    - manually:  python -m services.vector_service [--rebuild | --search "..."]
"""

from __future__ import annotations

import hashlib
import re
import threading
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence, Union

import chromadb
from chromadb.config import Settings

from core.config import Config
from core.logger import get_logger
from services.memory_service import MemoryService, Message
from services.ollama_service import OllamaService

logger = get_logger(__name__)

# ----------------------------------------------------------------------
# Constants
# ----------------------------------------------------------------------

CONVERSATIONS_COLLECTION = "conversations"
DOCUMENTS_COLLECTION = "documents"

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024  # skip larger files (logged)
UNCATEGORIZED = "uncategorized"          # category for files directly in data/documents/

CHUNK_SIZE = 1000         # target characters per document chunk
CHUNK_OVERLAP = 150       # characters carried over between neighbouring chunks
EMBED_BATCH_SIZE = 16     # texts per embedding request
MESSAGE_EMBED_MAX_CHARS = 4000  # long messages are truncated for embedding only

DEFAULT_TOP_K = 3
SNIPPET_MAX_CHARS = 600   # max characters per hit injected into a prompt
REQUEST_SYNC_BUDGET_SECONDS = 5.0

STATE_KEY_LAST_MESSAGE_ID = "vector_service.last_indexed_message_id"

# Qwen3-Embedding performs best when queries (not stored texts) carry an instruction.
QUERY_INSTRUCTION = (
    "Instruct: Given a question or statement, retrieve relevant past "
    "conversations, notes and document passages\nQuery: "
)

NO_BACKGROUND = "(Ingen relevant bakgrund hittades.)"


# ----------------------------------------------------------------------
# Data model
# ----------------------------------------------------------------------

@dataclass(frozen=True)
class SearchHit:
    """A single long-term memory search result."""

    source: str                 # "conversation" or "document"
    text: str
    distance: float             # cosine distance, lower = more similar
    metadata: dict[str, Any]

    @property
    def label(self) -> str:
        """Human-readable origin of the hit, used in prompts and CLI output."""
        if self.source == "conversation":
            speaker = "User" if self.metadata.get("role") == "user" else "Agent"
            agent = self.metadata.get("agent_id") or "okänd agent"
            date = self.metadata.get("created_date", "?")
            return f"{date} · {agent} · {speaker}"
        page = self.metadata.get("page", 0)
        path = self.metadata.get("source_path", "?")
        return f"{path}, s. {page}" if page else path


class VectorServiceError(Exception):
    """Raised for VectorService-specific failures."""


# ----------------------------------------------------------------------
# Service
# ----------------------------------------------------------------------

class VectorService:
    """
    Long-term semantic memory: conversation history + personal documents.

    Thread-safe: syncs are serialized by a lock; a request-time sync never waits
    for another sync already in progress (it simply searches what is indexed).
    """

    def __init__(
        self,
        memory_service: Optional[MemoryService] = None,
        ollama_service: Optional[OllamaService] = None,
        persist_dir: Union[str, Path] = Config.VECTOR_STORE_DIR,
        documents_dir: Union[str, Path] = Config.DOCUMENTS_DIR,
        max_distance: float = Config.VECTOR_MAX_DISTANCE,
    ) -> None:
        self.memory = memory_service or MemoryService()
        self.ollama = ollama_service or OllamaService()
        self.embed_model = self.ollama.embed_model
        self.max_distance = max_distance

        self._persist_dir = Path(persist_dir)
        self._documents_dir = Path(documents_dir)
        self._persist_dir.mkdir(parents=True, exist_ok=True)
        self._documents_dir.mkdir(parents=True, exist_ok=True)

        self._sync_lock = threading.Lock()
        self._warned_paths: set[str] = set()
        self._empty_files: dict[str, tuple[int, int]] = {}

        # Telemetry off: Jarvis is local-first, nothing leaves the machine.
        self._client = chromadb.PersistentClient(
            path=str(self._persist_dir),
            settings=Settings(anonymized_telemetry=False),
        )
        self._conversations = self._open_collection(CONVERSATIONS_COLLECTION)
        self._documents = self._open_collection(DOCUMENTS_COLLECTION)
        self._doc_state = self._load_document_state()

        logger.info(
            "VectorService initialized | embed_model=%s | conversations=%s | "
            "document_chunks=%s | documents_dir=%s",
            self.embed_model,
            self._conversations.count(),
            self._documents.count(),
            self._documents_dir,
        )

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def build_context(
        self,
        query: str,
        exclude_message_ids: Iterable[int] = (),
        top_k: int = DEFAULT_TOP_K,
    ) -> str:
        """
        Return a formatted "relevant background" block for an agent prompt.

        Runs a short, non-blocking sync first, then searches both conversation
        history (all agents) and documents. Never raises: on any failure the
        agent gets NO_BACKGROUND and answers without long-term memory.
        """
        self.sync_all(time_budget=REQUEST_SYNC_BUDGET_SECONDS, blocking=False)

        try:
            query_vector = self._embed_query(query)
            conversation_hits = self._search_conversations(
                query_vector, top_k, set(exclude_message_ids)
            )
            document_hits = self._search_documents(query_vector, top_k)
        except Exception as exc:
            logger.warning("Long-term memory lookup failed, continuing without it: %s", exc)
            return NO_BACKGROUND

        logger.info(
            "Long-term memory lookup | conversation_hits=%s | document_hits=%s",
            len(conversation_hits),
            len(document_hits),
        )
        return format_background(conversation_hits, document_hits)

    def search_conversations(
        self,
        query: str,
        top_k: int = DEFAULT_TOP_K,
        exclude_message_ids: Iterable[int] = (),
    ) -> list[SearchHit]:
        """Semantic search over all conversation history (all agents)."""
        return self._search_conversations(
            self._embed_query(query), top_k, set(exclude_message_ids)
        )

    def search_documents(
        self,
        query: str,
        top_k: int = DEFAULT_TOP_K,
        category: Optional[str] = None,
    ) -> list[SearchHit]:
        """Semantic search over personal documents, optionally within one category."""
        return self._search_documents(self._embed_query(query), top_k, category)

    def sync_all(
        self, time_budget: Optional[float] = None, blocking: bool = True
    ) -> dict[str, int]:
        """
        Bring the vector index up to date with SQLite and data/documents/.

        :param time_budget: Stop starting new work after this many seconds
                            (None = run to completion).
        :param blocking:    If False and another sync is running, skip instead of waiting.
        :return: Counters describing what was done. Never raises.
        """
        if not self._sync_lock.acquire(blocking=blocking):
            logger.debug("Sync already in progress; skipping this one.")
            return {}

        try:
            deadline = time.monotonic() + time_budget if time_budget else None
            result: dict[str, int] = {}

            try:
                result.update(self._sync_conversations(deadline))
            except Exception as exc:
                logger.warning("Conversation sync failed (will retry next sync): %s", exc)

            try:
                result.update(self._sync_documents(deadline))
            except Exception as exc:
                logger.warning("Document sync failed (will retry next sync): %s", exc)

            if any(result.values()):
                logger.info("Long-term memory synced | %s", result)
            return result
        finally:
            self._sync_lock.release()

    def rebuild_index(self) -> dict[str, int]:
        """Wipe the vector index completely and rebuild it from SQLite + documents."""
        with self._sync_lock:
            logger.warning("Rebuilding long-term memory index from scratch...")
            for name in (CONVERSATIONS_COLLECTION, DOCUMENTS_COLLECTION):
                self._client.delete_collection(name)
            self.memory.set_state(STATE_KEY_LAST_MESSAGE_ID, 0)
            self._conversations = self._open_collection(CONVERSATIONS_COLLECTION)
            self._documents = self._open_collection(DOCUMENTS_COLLECTION)
            self._doc_state = {}
            self._empty_files = {}
        return self.sync_all()

    def stats(self) -> dict[str, Any]:
        """Return index sizes and configuration, for diagnostics."""
        return {
            "embed_model": self.embed_model,
            "indexed_messages": self._conversations.count(),
            "document_chunks": self._documents.count(),
            "documents": len(self._doc_state),
            "last_indexed_message_id": int(
                self.memory.get_state(STATE_KEY_LAST_MESSAGE_ID, 0) or 0
            ),
            "max_distance": self.max_distance,
        }

    # ------------------------------------------------------------------
    # Collections
    # ------------------------------------------------------------------

    def _open_collection(self, name: str):
        """
        Open (or create) a collection bound to the current embedding model.
        If it was built with a different model, it is dropped and rebuilt,
        because vectors from different models are not comparable.
        """
        metadata = {"hnsw:space": "cosine", "embed_model": self.embed_model}
        try:
            collection = self._client.get_collection(name=name, embedding_function=None)
        except Exception:
            # Not found (exception type differs between ChromaDB versions).
            return self._client.create_collection(
                name=name, metadata=metadata, embedding_function=None
            )

        existing_model = (collection.metadata or {}).get("embed_model")
        if existing_model != self.embed_model:
            logger.warning(
                "Collection '%s' was built with embedding model '%s' but '%s' is configured. "
                "Re-indexing from scratch.",
                name,
                existing_model,
                self.embed_model,
            )
            self._client.delete_collection(name)
            collection = self._client.create_collection(
                name=name, metadata=metadata, embedding_function=None
            )
            if name == CONVERSATIONS_COLLECTION:
                self.memory.set_state(STATE_KEY_LAST_MESSAGE_ID, 0)

        return collection

    # ------------------------------------------------------------------
    # Conversation sync
    # ------------------------------------------------------------------

    def _sync_conversations(self, deadline: Optional[float]) -> dict[str, int]:
        """
        Mirror SQLite into the vector index: index new messages and remove
        vectors whose message was deleted from SQLite (e.g. "forget this").
        """
        last_id = int(self.memory.get_state(STATE_KEY_LAST_MESSAGE_ID, 0) or 0)
        removed = 0

        # Deletions: more vectors than messages up to the pointer -> drop orphans.
        if last_id > 0 and self._conversations.count() > self.memory.count_messages(max_id=last_id):
            removed = self._remove_orphaned_messages()

        # Self-healing: index wiped but pointer kept -> start over.
        if last_id > 0 and self._conversations.count() == 0:
            logger.warning("Conversation index is empty; re-indexing all messages.")
            last_id = 0
            self.memory.set_state(STATE_KEY_LAST_MESSAGE_ID, 0)

        indexed = 0
        while not _deadline_passed(deadline):
            batch = self.memory.get_messages_after(last_id, limit=EMBED_BATCH_SIZE)
            if not batch:
                break

            vectors = self.ollama.embed([m.content[:MESSAGE_EMBED_MAX_CHARS] for m in batch])
            self._conversations.upsert(
                ids=[f"msg-{m.id}" for m in batch],
                embeddings=vectors,
                documents=[m.content for m in batch],
                metadatas=[_message_metadata(m) for m in batch],
            )

            last_id = batch[-1].id
            self.memory.set_state(STATE_KEY_LAST_MESSAGE_ID, last_id)
            indexed += len(batch)

        return {"messages_indexed": indexed, "messages_removed": removed}

    def _remove_orphaned_messages(self) -> int:
        """Delete vectors for messages that no longer exist in SQLite."""
        indexed_ids = set(self._conversations.get(include=[])["ids"])
        valid_ids = {f"msg-{message_id}" for message_id in self.memory.get_message_ids()}
        orphans = sorted(indexed_ids - valid_ids)
        if orphans:
            self._conversations.delete(ids=orphans)
            logger.info("Removed %s deleted message(s) from long-term memory.", len(orphans))
        return len(orphans)

    # ------------------------------------------------------------------
    # Document sync
    # ------------------------------------------------------------------

    def _load_document_state(self) -> dict[str, dict[str, Any]]:
        """Rebuild the per-file index state from chunk metadata stored in ChromaDB."""
        state: dict[str, dict[str, Any]] = {}
        if self._documents.count() == 0:
            return state

        result = self._documents.get(include=["metadatas"])
        for meta in result.get("metadatas") or []:
            path = meta.get("source_path")
            if path and path not in state:
                state[path] = {
                    "mtime_ns": int(meta.get("file_mtime_ns", 0)),
                    "size": int(meta.get("file_size", 0)),
                    "hash": meta.get("file_hash", ""),
                }
        return state

    def _scan_documents(self) -> dict[str, Path]:
        """Return all indexable files under data/documents/, keyed by relative POSIX path."""
        files: dict[str, Path] = {}
        for path in self._documents_dir.rglob("*"):
            if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue

            rel_parts = path.relative_to(self._documents_dir).parts
            # Skip hidden files/folders and Office lock files (~$report.docx).
            if any(part.startswith(".") for part in rel_parts) or path.name.startswith("~$"):
                continue

            rel = "/".join(rel_parts)
            if path.stat().st_size > MAX_FILE_SIZE_BYTES:
                if rel not in self._warned_paths:
                    logger.warning("Skipping '%s': larger than %s MB.", rel, MAX_FILE_SIZE_BYTES // 2**20)
                    self._warned_paths.add(rel)
                continue

            files[rel] = path
        return files

    def _sync_documents(self, deadline: Optional[float]) -> dict[str, int]:
        """Index new/changed documents and remove deleted ones."""
        on_disk = self._scan_documents()
        added = updated = removed = 0

        for rel in sorted(set(self._doc_state) - set(on_disk)):
            self._documents.delete(where={"source_path": rel})
            del self._doc_state[rel]
            removed += 1
            logger.info("Removed deleted document from long-term memory: %s", rel)

        for rel, path in sorted(on_disk.items()):
            if _deadline_passed(deadline):
                break

            stat = path.stat()
            fingerprint = (stat.st_mtime_ns, stat.st_size)
            known = self._doc_state.get(rel)

            if known and (known["mtime_ns"], known["size"]) == fingerprint:
                continue
            if self._empty_files.get(rel) == fingerprint:
                continue

            file_hash = _file_sha256(path)
            if known and known["hash"] == file_hash:
                # Touched but not changed: just refresh the stored fingerprint.
                self._refresh_fingerprint(rel, stat)
                continue

            chunk_count = self._index_file(rel, path, stat, file_hash)
            if chunk_count == 0:
                self._empty_files[rel] = fingerprint
                continue

            if known:
                updated += 1
            else:
                added += 1

        return {"documents_added": added, "documents_updated": updated, "documents_removed": removed}

    def _index_file(self, rel: str, path: Path, stat: Any, file_hash: str) -> int:
        """Extract, chunk, embed and store one file. Returns the number of chunks."""
        try:
            segments = _extract_text(path)
        except Exception as exc:
            logger.warning("Could not read '%s': %s", rel, exc)
            return 0

        chunk_texts: list[str] = []
        chunk_meta: list[dict[str, Any]] = []
        category = rel.split("/")[0] if "/" in rel else UNCATEGORIZED
        indexed_at = datetime.now(timezone.utc).isoformat()

        for page, text in segments:
            for chunk_index, chunk in enumerate(_chunk_text(text)):
                chunk_texts.append(chunk)
                chunk_meta.append({
                    "source_path": rel,
                    "category": category,
                    "filename": path.name,
                    "page": page,
                    "chunk_index": chunk_index,
                    "file_mtime_ns": stat.st_mtime_ns,
                    "file_size": stat.st_size,
                    "file_hash": file_hash,
                    "indexed_at": indexed_at,
                })

        if not chunk_texts:
            logger.warning(
                "No extractable text in '%s' (scanned PDF or empty file?). Skipping.", rel
            )
            return 0

        # Embed everything first: if embedding fails, the old chunks stay intact.
        vectors: list[list[float]] = []
        for start in range(0, len(chunk_texts), EMBED_BATCH_SIZE):
            vectors.extend(self.ollama.embed(chunk_texts[start:start + EMBED_BATCH_SIZE]))

        path_id = hashlib.sha1(rel.encode("utf-8")).hexdigest()[:16]
        ids = [f"doc-{path_id}-p{m['page']}-c{m['chunk_index']}" for m in chunk_meta]

        self._documents.delete(where={"source_path": rel})
        self._documents.add(ids=ids, embeddings=vectors, documents=chunk_texts, metadatas=chunk_meta)

        self._doc_state[rel] = {"mtime_ns": stat.st_mtime_ns, "size": stat.st_size, "hash": file_hash}
        self._empty_files.pop(rel, None)
        logger.info("Indexed document '%s' (%s chunks).", rel, len(chunk_texts))
        return len(chunk_texts)

    def _refresh_fingerprint(self, rel: str, stat: Any) -> None:
        """Update stored mtime/size for an unchanged file so it isn't re-hashed next time."""
        existing = self._documents.get(where={"source_path": rel}, include=["metadatas"])
        ids = existing.get("ids") or []
        if ids:
            metadatas = [
                {**meta, "file_mtime_ns": stat.st_mtime_ns, "file_size": stat.st_size}
                for meta in existing["metadatas"]
            ]
            self._documents.update(ids=ids, metadatas=metadatas)
        self._doc_state[rel]["mtime_ns"] = stat.st_mtime_ns
        self._doc_state[rel]["size"] = stat.st_size

    # ------------------------------------------------------------------
    # Search
    # ------------------------------------------------------------------

    def _embed_query(self, query: str) -> list[float]:
        return self.ollama.embed([QUERY_INSTRUCTION + query.strip()])[0]

    def _search_conversations(
        self, query_vector: list[float], top_k: int, exclude_ids: set[int]
    ) -> list[SearchHit]:
        hits = self._query(
            self._conversations, "conversation", query_vector, n_results=top_k + len(exclude_ids) + 5
        )
        hits = [h for h in hits if h.metadata.get("message_id") not in exclude_ids]
        return hits[:top_k]

    def _search_documents(
        self, query_vector: list[float], top_k: int, category: Optional[str] = None
    ) -> list[SearchHit]:
        where = {"category": category} if category else None
        return self._query(self._documents, "document", query_vector, n_results=top_k, where=where)

    def _query(
        self,
        collection,
        source: str,
        query_vector: list[float],
        n_results: int,
        where: Optional[dict[str, Any]] = None,
    ) -> list[SearchHit]:
        """Run a vector query and drop hits weaker than max_distance."""
        count = collection.count()
        if count == 0:
            return []

        result = collection.query(
            query_embeddings=[query_vector],
            n_results=min(n_results, count),
            where=where,
            include=["documents", "metadatas", "distances"],
        )

        hits = []
        for text, meta, distance in zip(
            result["documents"][0], result["metadatas"][0], result["distances"][0]
        ):
            logger.debug("Candidate %s hit | distance=%.3f", source, distance)
            if distance <= self.max_distance:
                hits.append(SearchHit(source=source, text=text, distance=distance, metadata=meta))
        return hits


# ----------------------------------------------------------------------
# Formatting
# ----------------------------------------------------------------------

def format_background(
    conversation_hits: Sequence[SearchHit], document_hits: Sequence[SearchHit]
) -> str:
    """Format search hits into a compact, labeled block for an agent prompt."""
    if not conversation_hits and not document_hits:
        return NO_BACKGROUND

    lines: list[str] = []
    if conversation_hits:
        lines.append("Från tidigare konversationer:")
        lines.extend(f"- [{hit.label}] {_snippet(hit.text)}" for hit in conversation_hits)
    if document_hits:
        if lines:
            lines.append("")
        lines.append("Från dina dokument:")
        lines.extend(f"- [{hit.label}] {_snippet(hit.text)}" for hit in document_hits)
    return "\n".join(lines)


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------

def _deadline_passed(deadline: Optional[float]) -> bool:
    return deadline is not None and time.monotonic() >= deadline


def _snippet(text: str, max_chars: int = SNIPPET_MAX_CHARS) -> str:
    compact = " ".join(text.split())
    return compact if len(compact) <= max_chars else compact[: max_chars - 1].rstrip() + "…"


def _message_metadata(message: Message) -> dict[str, Any]:
    """ChromaDB metadata must be str/int/float/bool - no None values."""
    return {
        "message_id": message.id,
        "session_id": message.session_id,
        "role": message.role,
        "agent_id": message.agent_id or "",
        "context": message.context or "",
        "created_at": message.created_at,
        "created_date": message.created_at[:10],
    }


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _extract_text(path: Path) -> list[tuple[int, str]]:
    """
    Extract text from a supported file.
    Returns (page_number, text) segments; page_number is 0 when not applicable.
    """
    suffix = path.suffix.lower()

    if suffix in (".txt", ".md"):
        return [(0, path.read_text(encoding="utf-8-sig", errors="replace"))]

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise VectorServiceError("password-protected PDF") from exc
        return [(number, page.extract_text() or "") for number, page in enumerate(reader.pages, start=1)]

    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return [(0, "\n".join(parts))]

    raise VectorServiceError(f"Unsupported file type: {suffix}")


def _chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """
    Split text into chunks of roughly `size` characters along paragraph
    boundaries, carrying `overlap` characters of context into the next chunk.
    """
    text = re.sub(r"[ \t]+\n", "\n", text.replace("\r\n", "\n"))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []

    # 1. Paragraphs, with oversized paragraphs split on word boundaries.
    pieces: list[str] = []
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = paragraph.strip()
        while len(paragraph) > size:
            cut = paragraph.rfind(" ", 0, size)
            if cut < size // 2:
                cut = size
            pieces.append(paragraph[:cut].strip())
            paragraph = paragraph[cut:].strip()
        if paragraph:
            pieces.append(paragraph)

    # 2. Pack pieces into chunks, with overlap between neighbours.
    chunks: list[str] = []
    current = ""
    for piece in pieces:
        if not current:
            current = piece
        elif len(current) + 2 + len(piece) <= size:
            current = f"{current}\n\n{piece}"
        else:
            chunks.append(current)
            tail = current[-overlap:]
            space = tail.find(" ")
            tail = tail[space + 1:] if space != -1 else ""
            current = f"{tail} … {piece}" if tail else piece
    if current:
        chunks.append(current)
    return chunks


# ----------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------

if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Jarvis long-term memory (VectorService)")
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--rebuild", action="store_true", help="wipe and rebuild the whole index")
    group.add_argument("--search", metavar="TEXT", help="test a search and show distances")
    group.add_argument("--stats", action="store_true", help="show index statistics only")
    args = parser.parse_args()

    service = VectorService()

    if args.stats:
        pass
    elif args.rebuild:
        print("Rebuilding index...")
        print(service.rebuild_index())
    elif args.search:
        service.sync_all()
        print(f'\nSearch: "{args.search}"  (max_distance={service.max_distance})\n')
        vector = service._embed_query(args.search)
        original = service.max_distance
        service.max_distance = 2.0  # show all candidates, mark which would pass
        for title, hits in (
            ("Conversations", service._search_conversations(vector, 5, set())),
            ("Documents", service._search_documents(vector, 5)),
        ):
            print(f"== {title} ==")
            for hit in hits:
                mark = "PASS" if hit.distance <= original else "drop"
                print(f"  [{mark}] {hit.distance:.3f}  [{hit.label}] {_snippet(hit.text, 120)}")
            if not hits:
                print("  (empty)")
        service.max_distance = original
    else:
        print("Syncing long-term memory...")
        print(service.sync_all())

    print("\nStats:", service.stats())
